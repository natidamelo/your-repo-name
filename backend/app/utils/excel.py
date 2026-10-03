import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Color Palette
NAVY_HEADER = "0F172A"
WHITE = "FFFFFF"
LIGHT_GRAY = "F8FAFC"
BORDER_GRAY = "CBD5E1"
DARK_TEXT = "1E293B"

# Shift Type Fills
FILL_WORK = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")     # Emerald light
FILL_OFF = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")      # Slate light
FILL_AM = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")       # Amber light
FILL_PM = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")       # Purple light
FILL_SUNDAY = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")   # Indigo light
FILL_CRITICAL = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid") # Rose light
FILL_WARNING = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")  # Yellow light
FILL_BACKUP = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")   # Orange/Amber light
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

def get_thin_border():
    thin = Side(border_style="thin", color=BORDER_GRAY)
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header_cell(cell, text):
    cell.value = text
    cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = get_thin_border()

def get_shift_tasks(s):
    """
    Extracts structured list of tasks for a shift, falling back to primary/secondary tasks.
    """
    tasks = []
    if hasattr(s, "tasks") and s.tasks:
        for t in s.tasks:
            if t.task_name:
                tasks.append({
                    "name": t.task_name.strip(),
                    "start": t.start_time or s.start_time or "",
                    "end": t.end_time or s.end_time or "",
                    "is_backup": bool(getattr(t, "is_backup", False)),
                    "notes": getattr(t, "notes", "") or ""
                })
    if not tasks and s and s.shift_type != "OFF":
        # Fallback to primary_task / secondary_task
        if s.primary_task:
            for pt in s.primary_task.split(","):
                pt = pt.strip()
                if pt:
                    tasks.append({
                        "name": pt,
                        "start": s.start_time or "08:00",
                        "end": s.end_time or "17:00",
                        "is_backup": False,
                        "notes": ""
                    })
        if s.secondary_task:
            for st in s.secondary_task.split(","):
                st = st.strip()
                if st and st not in [x["name"] for x in tasks]:
                    tasks.append({
                        "name": st,
                        "start": s.start_time or "08:00",
                        "end": s.end_time or "17:00",
                        "is_backup": False,
                        "notes": ""
                    })
    return tasks

def calculate_duration_str(start_str: str, end_str: str) -> str:
    try:
        if not start_str or not end_str:
            return "—"
        s_h, s_m = map(int, start_str.split(":"))
        e_h, e_m = map(int, end_str.split(":"))
        s_total = s_h * 60 + s_m
        e_total = e_h * 60 + e_m
        if e_total < s_total:
            e_total += 24 * 60 # crossed midnight
        diff_hours = (e_total - s_total) / 60.0
        return f"{diff_hours:.1f} hrs" if diff_hours % 1 != 0 else f"{int(diff_hours)} hrs"
    except Exception:
        return "—"

def create_schedule_excel_workbook(period, days, shifts_by_day, conflicts, coverage_by_day) -> io.BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    border = get_thin_border()
    sorted_days = sorted(days, key=lambda d: d.date)

    # -------------------------------------------------------------
    # SHEET 1: 14-Day Staff Roster (With All Tasks & Times)
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Staff Roster")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    total_cols = len(sorted_days) + 2
    last_col_letter = get_column_letter(total_cols)
    ws1.merge_cells(f"A1:{last_col_letter}1")
    title_cell = ws1["A1"]
    start_disp = sorted_days[0].date if sorted_days else period.start_date
    end_disp = sorted_days[-1].date if sorted_days else period.end_date
    title_cell.value = f"CALL CENTER STAFF ROSTER — {period.name} ({start_disp} to {end_disp})"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    # Headers: Employee Name, Position, Day 1 to Day N
    ws1.row_dimensions[2].height = 28
    style_header_cell(ws1["A2"], "Employee Name")
    style_header_cell(ws1["B2"], "Position")

    for col_idx, day in enumerate(sorted_days, start=3):
        col_letter = get_column_letter(col_idx)
        d_obj = datetime.strptime(day.date, "%Y-%m-%d")
        header_text = f"{d_obj.strftime('%a')}\n{d_obj.strftime('%b %d')}"
        style_header_cell(ws1[f"{col_letter}2"], header_text)

    # Group shifts by employee
    emp_map = {}
    for day in sorted_days:
        for s in shifts_by_day.get(day.id, []):
            emp_name = s.employee.first_name if (hasattr(s, "employee") and s.employee) else getattr(s, "employee_name", "Staff")
            pos = s.employee.position if (hasattr(s, "employee") and s.employee and s.employee.position) else "Agent"
            if emp_name not in emp_map:
                emp_map[emp_name] = {"position": pos, "shifts": {}}
            emp_map[emp_name]["shifts"][day.date] = s

    row_num = 3
    for emp_name, info in sorted(emp_map.items()):
        cell_name = ws1[f"A{row_num}"]
        cell_name.value = emp_name
        cell_name.font = Font(name="Calibri", size=11, bold=True)
        cell_name.border = border
        cell_name.alignment = Alignment(vertical="center", horizontal="left")

        cell_pos = ws1[f"B{row_num}"]
        cell_pos.value = info["position"]
        cell_pos.font = Font(name="Calibri", size=10)
        cell_pos.border = border
        cell_pos.alignment = Alignment(vertical="center", horizontal="left")

        max_lines_in_row = 1

        for col_idx, day in enumerate(sorted_days, start=3):
            col_letter = get_column_letter(col_idx)
            cell = ws1[f"{col_letter}{row_num}"]
            cell.border = border
            s = info["shifts"].get(day.date)
            st = s.shift_type if s else "OFF"

            if st == "OFF" or not s:
                cell.value = "OFF\nRest Day"
                cell.font = Font(name="Calibri", size=9, bold=False, color="64748B")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.fill = FILL_OFF
            else:
                type_display = {
                    "WORK": "WORK",
                    "AM_HALF": "AM HALF",
                    "PM_HALF": "PM HALF",
                    "SUNDAY_DUTY": "SUNDAY"
                }.get(st, st)

                hours_header = f"{s.start_time}–{s.end_time}" if (s.start_time and s.end_time) else ""
                lunch_info = f" [Lunch: {s.lunch_start}–{s.lunch_end}]" if (s.lunch_start and s.lunch_end) else ""
                header_line = f"[{type_display} {hours_header}]{lunch_info}".strip() if hours_header else f"[{type_display}]{lunch_info}"

                tasks = get_shift_tasks(s)
                task_lines = []
                for t in tasks:
                    time_part = f"({t['start']}–{t['end']})" if (t['start'] and t['end']) else ""
                    if t["is_backup"]:
                        task_lines.append(f"• {t['name']} [BKP] {time_part}".strip())
                    else:
                        task_lines.append(f"• {t['name']} {time_part}".strip())

                if not task_lines:
                    task_lines.append("• Call Center")

                cell_text = header_line + "\n" + "\n".join(task_lines)
                cell.value = cell_text
                cell.font = Font(name="Calibri", size=9, bold=False)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

                if st == "WORK":
                    cell.fill = FILL_WORK
                elif st == "AM_HALF":
                    cell.fill = FILL_AM
                elif st == "PM_HALF":
                    cell.fill = FILL_PM
                elif st == "SUNDAY_DUTY":
                    cell.fill = FILL_SUNDAY

                line_count = len(cell_text.split("\n"))
                if line_count > max_lines_in_row:
                    max_lines_in_row = line_count

        # Dynamic row height so every task line is visible
        ws1.row_dimensions[row_num].height = max(34, 16 + max_lines_in_row * 14)
        row_num += 1

    # Column widths for Sheet 1
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 20
    for col_idx in range(3, len(sorted_days) + 3):
        ws1.column_dimensions[get_column_letter(col_idx)].width = 28

    # Legend at the bottom of Sheet 1
    legend_row = row_num + 1
    ws1.merge_cells(f"A{legend_row}:{last_col_letter}{legend_row}")
    l_cell = ws1[f"A{legend_row}"]
    l_cell.value = (
        "Legend: [WORK] Full Shift  |  [OFF] Day Off  |  [AM HALF] Saturday Morning  |  "
        "[PM HALF] Saturday Afternoon  |  [SUNDAY] Operating Squad  |  [BKP] Standby / Afternoon Backup Slot"
    )
    l_cell.font = Font(name="Calibri", size=9, italic=True, color="475569")
    l_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[legend_row].height = 24

    # -------------------------------------------------------------
    # SHEET 2: Daily Call-Center Coverage
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Daily Coverage")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    t2 = ws2["A1"]
    t2.value = "DAILY CALL-CENTER COVERAGE BREAKDOWN"
    t2.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    t2.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 34

    headers_cov = ["Date", "Day", "8:00–9:00", "9:00–12:00", "12:00–1:00", "1:00–2:00", "2:00–5:00", "5:00–6:00"]
    ws2.row_dimensions[2].height = 26
    for idx, h in enumerate(headers_cov, start=1):
        style_header_cell(ws2[f"{get_column_letter(idx)}2"], h)

    r_cov = 3
    for day in sorted_days:
        ws2.row_dimensions[r_cov].height = 22
        d_obj = datetime.strptime(day.date, "%Y-%m-%d")
        ws2[f"A{r_cov}"].value = day.date
        ws2[f"A{r_cov}"].border = border
        ws2[f"A{r_cov}"].alignment = Alignment(horizontal="center", vertical="center")

        ws2[f"B{r_cov}"].value = d_obj.strftime("%A")
        ws2[f"B{r_cov}"].border = border
        ws2[f"B{r_cov}"].alignment = Alignment(horizontal="center", vertical="center")

        cov_data = coverage_by_day.get(day.date, {})
        slots = cov_data.get("time_slots", [])
        for s_idx, slot in enumerate(slots, start=3):
            cell = ws2[f"{get_column_letter(s_idx)}{r_cov}"]
            cell.value = f"{slot.get('staff_count', 0)} staff ({slot.get('status', 'RED')})"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            status = slot.get("status")
            if status == "GREEN":
                cell.fill = FILL_WORK
            elif status == "YELLOW":
                cell.fill = FILL_WARNING
            else:
                cell.fill = FILL_CRITICAL
        r_cov += 1

    for c in range(1, 9):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # -------------------------------------------------------------
    # SHEET 3: Sunday Schedule (Operating Squad)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Sunday Schedule")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:F1")
    t3 = ws3["A1"]
    t3.value = "SUNDAY OPERATING SQUAD SCHEDULE (EXACTLY 4 STAFF)"
    t3.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    t3.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t3.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 34

    headers_sun = ["Sunday Date", "Staff Member", "Position", "Shift Hours", "Lunch Break", "All Assigned Tasks & Times"]
    ws3.row_dimensions[2].height = 26
    for idx, h in enumerate(headers_sun, start=1):
        style_header_cell(ws3[f"{get_column_letter(idx)}2"], h)

    r_sun = 3
    for day in sorted_days:
        d_obj = datetime.strptime(day.date, "%Y-%m-%d")
        if d_obj.weekday() == 6:
            day_shifts = shifts_by_day.get(day.id, [])
            working = [s for s in day_shifts if s.shift_type in ("WORK", "SUNDAY_DUTY", "AM_HALF", "PM_HALF")]
            for s in working:
                emp_name = s.employee.first_name if (hasattr(s, "employee") and s.employee) else getattr(s, "employee_name", "Staff")
                pos = s.employee.position if (hasattr(s, "employee") and s.employee and s.employee.position) else "Agent"
                s_tasks = get_shift_tasks(s)
                task_strs = []
                for t in s_tasks:
                    time_part = f"({t['start']}–{t['end']})" if t['start'] and t['end'] else ""
                    bkp_part = " [BKP]" if t["is_backup"] else ""
                    task_strs.append(f"• {t['name']}{bkp_part} {time_part}".strip())

                tasks_val = "\n".join(task_strs) if task_strs else "• Call Center"

                ws3.row_dimensions[r_sun].height = max(24, 14 + len(task_strs) * 14)
                ws3[f"A{r_sun}"].value = f"{day.date} ({d_obj.strftime('%B %d, %Y')})"
                ws3[f"B{r_sun}"].value = emp_name
                ws3[f"C{r_sun}"].value = pos
                ws3[f"D{r_sun}"].value = f"{s.start_time} – {s.end_time}" if s.start_time else "—"
                ws3[f"E{r_sun}"].value = f"{s.lunch_start} – {s.lunch_end}" if s.lunch_start else "—"
                ws3[f"F{r_sun}"].value = tasks_val

                for c in range(1, 7):
                    ws3[f"{get_column_letter(c)}{r_sun}"].border = border
                    ws3[f"{get_column_letter(c)}{r_sun}"].alignment = Alignment(vertical="center", wrap_text=True)

                ws3[f"B{r_sun}"].font = Font(bold=True)
                ws3[f"B{r_sun}"].fill = FILL_SUNDAY
                r_sun += 1

    if r_sun == 3:
        ws3.merge_cells("A3:F3")
        c = ws3["A3"]
        c.value = "No Sundays in the selected date range."
        c.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws3.row_dimensions[3].height = 24
        for col_idx in range(1, 7):
            ws3[f"{get_column_letter(col_idx)}3"].border = border

    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 20
    ws3.column_dimensions["C"].width = 18
    ws3.column_dimensions["D"].width = 18
    ws3.column_dimensions["E"].width = 18
    ws3.column_dimensions["F"].width = 38

    # -------------------------------------------------------------
    # SHEET 4: Staff Assignments (Master Table with All Tasks)
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Staff Assignments")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:J1")
    t4 = ws4["A1"]
    t4.value = "DAILY STAFF TASK & SHIFT ASSIGNMENTS (ALL TASKS & TIME SLOTS)"
    t4.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    t4.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t4.alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 34

    headers_asgn = [
        "Date", "Day", "Employee Name", "Position", "Shift Type",
        "Shift Hours", "Lunch Break", "Assigned Tasks & Hours", "Backup Coverage", "Notes"
    ]
    ws4.row_dimensions[2].height = 26
    for idx, h in enumerate(headers_asgn, start=1):
        style_header_cell(ws4[f"{get_column_letter(idx)}2"], h)

    r_asgn = 3
    for day in sorted_days:
        d_obj = datetime.strptime(day.date, "%Y-%m-%d")
        for s in shifts_by_day.get(day.id, []):
            emp_name = s.employee.first_name if (hasattr(s, "employee") and s.employee) else getattr(s, "employee_name", "Staff")
            pos = s.employee.position if (hasattr(s, "employee") and s.employee and s.employee.position) else "Agent"

            all_tasks = get_shift_tasks(s)
            regular_tasks = [t for t in all_tasks if not t["is_backup"]]
            backup_tasks = [t for t in all_tasks if t["is_backup"]]

            reg_str = "\n".join([f"• {t['name']} ({t['start']}–{t['end']})" for t in regular_tasks]) if regular_tasks else ("—" if s.shift_type == "OFF" else "• Call Center")
            bkp_str = "\n".join([f"• {t['name']} [BKP] ({t['start']}–{t['end']})" for t in backup_tasks]) if backup_tasks else "—"

            max_items = max(len(regular_tasks), len(backup_tasks), 1)
            ws4.row_dimensions[r_asgn].height = max(22, 12 + max_items * 14)

            ws4[f"A{r_asgn}"].value = day.date
            ws4[f"B{r_asgn}"].value = d_obj.strftime("%A")
            ws4[f"C{r_asgn}"].value = emp_name
            ws4[f"D{r_asgn}"].value = pos
            ws4[f"E{r_asgn}"].value = s.shift_type
            ws4[f"F{r_asgn}"].value = f"{s.start_time} – {s.end_time}" if s.start_time else "—"
            ws4[f"G{r_asgn}"].value = f"{s.lunch_start} – {s.lunch_end}" if s.lunch_start else "—"
            ws4[f"H{r_asgn}"].value = reg_str
            ws4[f"I{r_asgn}"].value = bkp_str
            ws4[f"J{r_asgn}"].value = s.notes or "—"

            for c in range(1, 11):
                cell = ws4[f"{get_column_letter(c)}{r_asgn}"]
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

            if s.shift_type == "OFF":
                ws4[f"E{r_asgn}"].fill = FILL_OFF
            elif s.shift_type == "WORK":
                ws4[f"E{r_asgn}"].fill = FILL_WORK
            elif s.shift_type == "AM_HALF":
                ws4[f"E{r_asgn}"].fill = FILL_AM
            elif s.shift_type == "PM_HALF":
                ws4[f"E{r_asgn}"].fill = FILL_PM
            elif s.shift_type == "SUNDAY_DUTY":
                ws4[f"E{r_asgn}"].fill = FILL_SUNDAY

            if backup_tasks:
                ws4[f"I{r_asgn}"].fill = FILL_BACKUP

            r_asgn += 1

    if r_asgn == 3:
        ws4.merge_cells("A3:J3")
        c = ws4["A3"]
        c.value = "No staff assignments in the selected date range."
        c.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws4.row_dimensions[3].height = 24
        for col_idx in range(1, 11):
            ws4[f"{get_column_letter(col_idx)}3"].border = border

    ws4.column_dimensions["A"].width = 13
    ws4.column_dimensions["B"].width = 13
    ws4.column_dimensions["C"].width = 18
    ws4.column_dimensions["D"].width = 18
    ws4.column_dimensions["E"].width = 14
    ws4.column_dimensions["F"].width = 16
    ws4.column_dimensions["G"].width = 16
    ws4.column_dimensions["H"].width = 36
    ws4.column_dimensions["I"].width = 30
    ws4.column_dimensions["J"].width = 22

    # -------------------------------------------------------------
    # -------------------------------------------------------------
    # SHEET 5: Task Schedule & Coverage Role Breakdown
    # -------------------------------------------------------------
    ws_tasks = wb.create_sheet(title="Task Breakdown")
    ws_tasks.views.sheetView[0].showGridLines = True
    ws_tasks.freeze_panes = "A5"

    # Color Palette & Styles for Channels
    channel_styles = {
        "Call Center": {
            "fill": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="1E40AF")
        },
        "Telegram": {
            "fill": PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="0F766E")
        },
        "GDS": {
            "fill": PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="6D28D9")
        },
        "Amadeus": {
            "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="B45309")
        },
        "2839 phone": {
            "fill": PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="9D174D")
        },
        "Email": {
            "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="334155")
        },
        "ELMS": {
            "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="334155")
        },
        "QUE": {
            "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "font": Font(name="Calibri", size=10, bold=True, color="334155")
        }
    }
    def_channel_style = {
        "fill": PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=True, color="1E293B")
    }

    # Role Category Badges
    role_category_styles = {
        "CORE": {
            "fill": PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="4338CA")
        },
        "OPENING": {
            "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="15803D")
        },
        "LUNCH COVER": {
            "fill": PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="9F1239")
        },
        "BACKUP (BKP)": {
            "fill": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="C2410C")
        },
        "SECONDARY": {
            "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="475569")
        },
        "EVENING": {
            "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="B45309")
        },
        "REGULAR": {
            "fill": PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"),
            "font": Font(name="Calibri", size=9, bold=True, color="334155")
        }
    }

    def get_role_meta(task, shift):
        t_name = task.get("name", "").strip()
        is_bkp = bool(task.get("is_backup", False))
        t_start = task.get("start") or shift.start_time or ""
        t_end = task.get("end") or shift.end_time or ""
        emp_pos = shift.employee.position if (hasattr(shift, "employee") and shift.employee and shift.employee.position) else ""

        # 1. Explicit Backup
        if is_bkp:
            return "Standby / Backup Slot (BKP)", "BACKUP (BKP)"

        # 2. Specific Time Slot Windows
        if t_start == "08:00" and t_end == "09:00":
            return "Morning Opening Cover (08:00–09:00)", "OPENING"
        if (t_start in ("12:00", "13:00") and t_end in ("13:00", "14:00")) or ("lunch" in t_name.lower()):
            return "Lunch Interval Cover (12:00–14:00)", "LUNCH COVER"
        if t_start == "17:00" and t_end == "18:00":
            return "Evening Handover Cover (17:00–18:00)", "EVENING"

        # 3. Sunday Squad Duty
        if shift.shift_type == "SUNDAY_DUTY":
            if t_name == (shift.primary_task or "Call Center"):
                return "Sunday Squad Primary Duty", "CORE"
            return "Sunday Squad Support Channel", "SECONDARY"

        # 4. Position & Special Rule Context
        if "Backup Telegram" in emp_pos and t_name == "Telegram":
            return "Telegram Standby & Overflow Cover", "BACKUP (BKP)"
        if "Lunch CC Cover" in emp_pos and t_name == "Call Center":
            return "Lunch Peak Voice Support", "LUNCH COVER"
        if "Beti Coverage Lead" in emp_pos and t_name == "GDS":
            return "GDS Escalation & Coverage Lead", "CORE"

        # 5. Primary vs Secondary channel
        prim_tasks = [p.strip() for p in (shift.primary_task or "").split(",") if p.strip()]
        sec_tasks = [s.strip() for s in (shift.secondary_task or "").split(",") if s.strip()]

        if t_name in prim_tasks:
            if t_start == shift.start_time and t_end == shift.end_time:
                return "Primary Core Channel (Full Shift)", "CORE"
            return "Primary Channel Core Duty", "CORE"
        elif t_name in sec_tasks:
            return "Secondary Channel Support", "SECONDARY"
        elif t_start == shift.start_time and t_end == shift.end_time:
            return "Full-Shift Channel Duty", "CORE"
        else:
            return "Scheduled Channel Assignment", "REGULAR"

    # Row 1: Title Header Banner
    ws_tasks.merge_cells("A1:L1")
    tt = ws_tasks["A1"]
    tt.value = f"CALL CENTER TASK BREAKDOWN & COVERAGE ROLE SCHEDULE — {period.name}"
    tt.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    tt.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    tt.alignment = Alignment(horizontal="center", vertical="center")
    ws_tasks.row_dimensions[1].height = 36

    # Row 2: Sub-banner
    ws_tasks.merge_cells("A2:L2")
    sub_title = ws_tasks["A2"]
    start_disp = sorted_days[0].date if sorted_days else period.start_date
    end_disp = sorted_days[-1].date if sorted_days else period.end_date
    sub_title.value = f"Period: {start_disp} to {end_disp}   •   Organized by Date & Staff Member   •   Color-Coded Channels & Operational Coverage Roles"
    sub_title.font = Font(name="Calibri", size=10, italic=True, bold=True, color="334155")
    sub_title.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    sub_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_tasks.row_dimensions[2].height = 22

    # Row 3: Quick Key / Legend
    ws_tasks.merge_cells("A3:L3")
    leg_cell = ws_tasks["A3"]
    leg_cell.value = (
        "Channel Keys: [Call Center: Blue]  [Telegram: Teal]  [GDS: Purple]  [Amadeus: Amber]  [2839 Phone: Pink]  |  "
        "Coverage Roles: [CORE] Full-Shift Duty  |  [OPENING] 08:00–09:00  |  [LUNCH COVER] 12:00–14:00  |  [BACKUP] Standby  |  [SECONDARY] Cross-Skill"
    )
    leg_cell.font = Font(name="Calibri", size=9, italic=True, color="475569")
    leg_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    leg_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_tasks.row_dimensions[3].height = 20

    # Row 4: Column Headers
    headers_tb = [
        "Date", "Day", "Staff Member", "Position", "Shift Hours",
        "Assigned Channel", "Coverage Slot", "Duration",
        "Coverage Role Description", "Role Category", "Shift Type", "Staff Notes & Rules"
    ]
    ws_tasks.row_dimensions[4].height = 28
    for idx, h in enumerate(headers_tb, start=1):
        style_header_cell(ws_tasks[f"{get_column_letter(idx)}4"], h)

    r_tb = 5
    emp_palette = [
        PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
        PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    ]

    for day_idx, day in enumerate(sorted_days, start=1):
        d_obj = datetime.strptime(day.date, "%Y-%m-%d")
        day_shifts = shifts_by_day.get(day.id, [])
        working_shifts = [s for s in day_shifts if s.shift_type != "OFF"]

        if not working_shifts:
            continue

        # Count daily metrics
        day_tasks_total = 0
        day_backups_total = 0
        for s in working_shifts:
            t_list = get_shift_tasks(s)
            day_tasks_total += max(len(t_list), 1)
            day_backups_total += sum(1 for t in t_list if t.get("is_backup"))

        # Day Section Header Bar
        ws_tasks.row_dimensions[r_tb].height = 26
        ws_tasks.merge_cells(f"A{r_tb}:L{r_tb}")
        day_bar = ws_tasks[f"A{r_tb}"]
        day_bar.value = (
            f"DAY {day_idx}: {d_obj.strftime('%A, %B %d, %Y').upper()}   •   "
            f"{len(working_shifts)} Active Staff On Duty   •   "
            f"{day_tasks_total} Channel Slots Allocated   •   "
            f"{day_backups_total} Standby Backup Roles"
        )
        day_bar.font = Font(name="Calibri", size=10.5, bold=True, color=WHITE)
        day_bar.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        day_bar.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        for c in range(1, 13):
            ws_tasks[f"{get_column_letter(c)}{r_tb}"].border = border
        r_tb += 1

        for emp_idx, s in enumerate(working_shifts):
            emp_name = s.employee.first_name if (hasattr(s, "employee") and s.employee) else getattr(s, "employee_name", "Staff")
            pos = s.employee.position if (hasattr(s, "employee") and s.employee and s.employee.position) else "Agent"
            tasks = get_shift_tasks(s)
            if not tasks:
                tasks = [{"name": s.primary_task or "Call Center", "start": s.start_time or "08:00", "end": s.end_time or "17:00", "is_backup": False, "notes": ""}]

            emp_fill = emp_palette[emp_idx % 2]
            shift_hours_str = f"{s.start_time} – {s.end_time}" if s.start_time else "—"

            for t_idx, t in enumerate(tasks):
                dur_str = calculate_duration_str(t.get("start"), t.get("end"))
                role_label, role_badge = get_role_meta(t, s)
                task_name = t.get("name", "Call Center")
                time_slot_str = f"{t['start']} – {t['end']}" if (t.get("start") and t.get("end")) else shift_hours_str

                ws_tasks.row_dimensions[r_tb].height = 23

                # Col A: Date
                ws_tasks[f"A{r_tb}"].value = day.date
                ws_tasks[f"A{r_tb}"].font = Font(name="Calibri", size=10, color="0F172A" if t_idx == 0 else "94A3B8", bold=(t_idx == 0))
                ws_tasks[f"A{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"A{r_tb}"].fill = emp_fill

                # Col B: Day
                ws_tasks[f"B{r_tb}"].value = d_obj.strftime("%A")
                ws_tasks[f"B{r_tb}"].font = Font(name="Calibri", size=10, color="0F172A" if t_idx == 0 else "94A3B8", bold=(t_idx == 0))
                ws_tasks[f"B{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"B{r_tb}"].fill = emp_fill

                # Col C: Staff Member
                emp_display = emp_name if t_idx == 0 else f"  ↳ {emp_name}"
                ws_tasks[f"C{r_tb}"].value = emp_display
                ws_tasks[f"C{r_tb}"].font = Font(name="Calibri", size=11 if t_idx == 0 else 9.5, bold=True, color="0F172A" if t_idx == 0 else "64748B")
                ws_tasks[f"C{r_tb}"].alignment = Alignment(horizontal="left", vertical="center", indent=0 if t_idx == 0 else 1)
                ws_tasks[f"C{r_tb}"].fill = emp_fill

                # Col D: Position
                ws_tasks[f"D{r_tb}"].value = pos if t_idx == 0 else "—"
                ws_tasks[f"D{r_tb}"].font = Font(name="Calibri", size=9.5 if t_idx == 0 else 9, color="475569" if t_idx == 0 else "94A3B8")
                ws_tasks[f"D{r_tb}"].alignment = Alignment(horizontal="left", vertical="center")
                ws_tasks[f"D{r_tb}"].fill = emp_fill

                # Col E: Shift Hours
                ws_tasks[f"E{r_tb}"].value = shift_hours_str
                ws_tasks[f"E{r_tb}"].font = Font(name="Calibri", size=10 if t_idx == 0 else 9, bold=(t_idx == 0), color="0F172A" if t_idx == 0 else "94A3B8")
                ws_tasks[f"E{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"E{r_tb}"].fill = emp_fill

                # Col F: Assigned Channel
                t_style = channel_styles.get(task_name, def_channel_style)
                ws_tasks[f"F{r_tb}"].value = f"● {task_name}"
                ws_tasks[f"F{r_tb}"].font = t_style["font"]
                ws_tasks[f"F{r_tb}"].fill = t_style["fill"]
                ws_tasks[f"F{r_tb}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

                # Col G: Coverage Slot
                ws_tasks[f"G{r_tb}"].value = time_slot_str
                ws_tasks[f"G{r_tb}"].font = Font(name="Calibri", size=10, bold=True, color="1E293B")
                ws_tasks[f"G{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"G{r_tb}"].fill = emp_fill

                # Col H: Duration
                ws_tasks[f"H{r_tb}"].value = dur_str
                ws_tasks[f"H{r_tb}"].font = Font(name="Calibri", size=10, color="475569")
                ws_tasks[f"H{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"H{r_tb}"].fill = emp_fill

                # Col I: Coverage Role Description
                ws_tasks[f"I{r_tb}"].value = role_label
                ws_tasks[f"I{r_tb}"].font = Font(name="Calibri", size=9.5, bold=(role_badge in ("CORE", "OPENING", "BACKUP (BKP)")), color="0F172A" if t_idx == 0 else "334155")
                ws_tasks[f"I{r_tb}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
                ws_tasks[f"I{r_tb}"].fill = emp_fill

                # Col J: Role Category Badge
                badge_style = role_category_styles.get(role_badge, role_category_styles["REGULAR"])
                ws_tasks[f"J{r_tb}"].value = role_badge
                ws_tasks[f"J{r_tb}"].font = badge_style["font"]
                ws_tasks[f"J{r_tb}"].fill = badge_style["fill"]
                ws_tasks[f"J{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")

                # Col K: Shift Type
                st = s.shift_type
                ws_tasks[f"K{r_tb}"].value = st if t_idx == 0 else st
                ws_tasks[f"K{r_tb}"].alignment = Alignment(horizontal="center", vertical="center")
                ws_tasks[f"K{r_tb}"].font = Font(name="Calibri", size=9, bold=True, color="0F172A" if t_idx == 0 else "94A3B8")
                if t_idx == 0:
                    if st == "WORK":
                        ws_tasks[f"K{r_tb}"].fill = FILL_WORK
                    elif st == "AM_HALF":
                        ws_tasks[f"K{r_tb}"].fill = FILL_AM
                    elif st == "PM_HALF":
                        ws_tasks[f"K{r_tb}"].fill = FILL_PM
                    elif st == "SUNDAY_DUTY":
                        ws_tasks[f"K{r_tb}"].fill = FILL_SUNDAY
                    else:
                        ws_tasks[f"K{r_tb}"].fill = emp_fill
                else:
                    ws_tasks[f"K{r_tb}"].fill = emp_fill

                # Col L: Staff Notes & Rules
                note_val = (s.notes or "") if t_idx == 0 else (t.get("notes") or "")
                ws_tasks[f"L{r_tb}"].value = note_val or "—"
                ws_tasks[f"L{r_tb}"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
                ws_tasks[f"L{r_tb}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
                ws_tasks[f"L{r_tb}"].fill = emp_fill

                # Grid borders
                for c in range(1, 13):
                    ws_tasks[f"{get_column_letter(c)}{r_tb}"].border = border

                r_tb += 1

    if r_tb == 5:
        ws_tasks.merge_cells("A5:L5")
        c = ws_tasks["A5"]
        c.value = "No task assignments in the selected date range."
        c.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws_tasks.row_dimensions[5].height = 24
        for col_idx in range(1, 13):
            ws_tasks[f"{get_column_letter(col_idx)}5"].border = border
        last_row = 5
    else:
        last_row = r_tb - 1

    # AutoFilter on headers
    ws_tasks.auto_filter.ref = f"A4:L{last_row}"

    # Column widths for Sheet 5
    ws_tasks.column_dimensions["A"].width = 13
    ws_tasks.column_dimensions["B"].width = 12
    ws_tasks.column_dimensions["C"].width = 20
    ws_tasks.column_dimensions["D"].width = 26
    ws_tasks.column_dimensions["E"].width = 16
    ws_tasks.column_dimensions["F"].width = 18
    ws_tasks.column_dimensions["G"].width = 18
    ws_tasks.column_dimensions["H"].width = 11
    ws_tasks.column_dimensions["I"].width = 34
    ws_tasks.column_dimensions["J"].width = 16
    ws_tasks.column_dimensions["K"].width = 14
    ws_tasks.column_dimensions["L"].width = 26

    # -------------------------------------------------------------
    # SHEET 6: Conflict & Validation Audit Report
    # -------------------------------------------------------------
    ws6 = wb.create_sheet(title="Conflict Report")
    ws6.views.sheetView[0].showGridLines = True

    ws6.merge_cells("A1:E1")
    t6 = ws6["A1"]
    t6.value = "SCHEDULE VALIDATION & CONFLICT AUDIT REPORT"
    t6.font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    t6.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t6.alignment = Alignment(horizontal="center", vertical="center")
    ws6.row_dimensions[1].height = 34

    headers_conf = ["Date / Target", "Employee", "Severity", "Error / Warning Type", "Detailed Message & Suggestion"]
    ws6.row_dimensions[2].height = 26
    for idx, h in enumerate(headers_conf, start=1):
        style_header_cell(ws6[f"{get_column_letter(idx)}2"], h)

    r_conf = 3
    if not conflicts:
        ws6.row_dimensions[r_conf].height = 26
        ws6.merge_cells(f"A{r_conf}:E{r_conf}")
        c = ws6[f"A{r_conf}"]
        c.value = "✓ VALIDATION PASSED: Zero conflicts detected. All constraints strictly satisfied."
        c.font = Font(name="Calibri", size=11, bold=True, color="047857")
        c.fill = FILL_WORK
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    else:
        for item in conflicts:
            ws6.row_dimensions[r_conf].height = 24
            c_date = item.date if hasattr(item, "date") else item.get("date", "—")
            c_emp = item.employee_name if hasattr(item, "employee_name") else item.get("employee_name", "—")
            c_sev = item.severity if hasattr(item, "severity") else item.get("severity", "critical")
            c_type = item.error_type if hasattr(item, "error_type") else item.get("error_type", "—")
            c_msg = item.message if hasattr(item, "message") else item.get("message", "")
            c_sug = item.suggestion if hasattr(item, "suggestion") else item.get("suggestion", "")

            ws6[f"A{r_conf}"].value = c_date
            ws6[f"B{r_conf}"].value = c_emp
            ws6[f"C{r_conf}"].value = (c_sev or "critical").upper()
            ws6[f"D{r_conf}"].value = c_type
            ws6[f"E{r_conf}"].value = f"{c_msg} (Suggestion: {c_sug})" if c_sug else c_msg

            for c in range(1, 6):
                cell = ws6[f"{get_column_letter(c)}{r_conf}"]
                cell.border = border
                cell.alignment = Alignment(vertical="center")

            if (c_sev or "").lower() == "critical":
                ws6[f"C{r_conf}"].fill = FILL_CRITICAL
                ws6[f"C{r_conf}"].font = Font(bold=True, color="991B1B")
            else:
                ws6[f"C{r_conf}"].fill = FILL_WARNING
                ws6[f"C{r_conf}"].font = Font(bold=True, color="854D0E")
            r_conf += 1

    ws6.column_dimensions["A"].width = 16
    ws6.column_dimensions["B"].width = 18
    ws6.column_dimensions["C"].width = 14
    ws6.column_dimensions["D"].width = 28
    ws6.column_dimensions["E"].width = 65

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
